"""crear_excel_ejemplo_cap06.py — Genera los Excel de ejemplo de la app «datos propios» del Cap. 6

Autor:       Dr. Aboud Barsekh-Onji · IPN — ESCA Unidad Santo Tomás
Descripción: Dos libros de ejemplo (y plantilla) para cap06_alternativas_datos_propios.html:
             «cap06_datos_ejemplo.xlsx»          — las seis alternativas ficticias del laboratorio del libro
                                                   (hojas «Alternativas» y «Criterios»).
             «cap06_datos_ejemplo_servicio.xlsx» — ocho servicios de entrega ficticios con cinco atributos.
             También escribe datos/cap06_ejemplo.json, que la app incluye para sus botones de ejemplo.
             Todos los datos son SINTÉTICOS.
Uso:         conda run -n research python matlab/utils/apps/datos/crear_excel_ejemplo_cap06.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).resolve().parent
CARPETA = AQUI.parents[2] / "cap06"
GUINDA = "6F1D46"

LIBRO = {
    "Alternativas": [
        ["Alternativa", "Precio_MXN", "Calidad_1a10", "Entrega_dias", "Resenas_1a5"],
        ["Alfa", 899, 6.0, 2, 4.1], ["Beta", 1499, 8.5, 5, 4.6], ["Gama", 1199, 7.5, 3, 4.3],
        ["Delta", 699, 4.5, 7, 3.6], ["Épsilon", 1799, 9.0, 4, 4.8], ["Zeta", 1049, 7.0, 2, 3.9],
    ],
    "Criterios": [
        ["Criterio", "Tipo", "Peso"],
        ["Precio_MXN", "costo", 0.40], ["Calidad_1a10", "beneficio", 0.30],
        ["Entrega_dias", "costo", 0.15], ["Resenas_1a5", "beneficio", 0.15],
    ],
}
SERVICIO = {
    "Alternativas": [
        ["Servicio", "Costo_envio_MXN", "Tiempo_horas", "Cobertura_pct", "Calificacion_1a5", "Seguimiento_1a10"],
        ["Mensajería Norte", 69, 48, 82, 4.2, 7], ["Express Sur", 129, 24, 70, 4.5, 9], ["Paquetería Centro", 89, 36, 95, 3.9, 6],
        ["Moto Urbana", 59, 6, 25, 4.4, 8], ["Logística Plus", 149, 24, 98, 4.7, 10], ["Envíos Económicos", 45, 72, 90, 3.5, 4],
        ["Correo Nacional", 79, 96, 100, 3.7, 5], ["Entrega Verde", 99, 48, 60, 4.3, 7],
    ],
    "Criterios": [
        ["Criterio", "Tipo", "Peso"],
        ["Costo_envio_MXN", "costo", 3], ["Tiempo_horas", "costo", 3], ["Cobertura_pct", "beneficio", 2],
        ["Calificacion_1a5", "beneficio", 2], ["Seguimiento_1a10", "beneficio", 1],
    ],
}


def estilo(ws, anchos):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=GUINDA); c.alignment = Alignment(horizontal="center")
    for col, ancho in zip("ABCDEFGH", anchos):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"


def libro(datos, ruta, titulo):
    wb = Workbook()
    ws = wb.active; ws.title = "Alternativas"
    for f in datos["Alternativas"]:
        ws.append(f)
    estilo(ws, (22, 17, 17, 16, 18, 18))
    wc = wb.create_sheet("Criterios")
    for f in datos["Criterios"]:
        wc.append(f)
    estilo(wc, (20, 12, 8))
    wb.properties.creator = "Instituto Politécnico Nacional - Dr. Aboud Barsekh Onji"; wb.properties.title = titulo
    CARPETA.mkdir(parents=True, exist_ok=True)
    wb.save(ruta); print(f"[OK] {ruta}")


def main():
    libro(LIBRO, CARPETA / "cap06_datos_ejemplo.xlsx", "Datos de ejemplo — alternativas del laboratorio (Cap. 6)")
    libro(SERVICIO, CARPETA / "cap06_datos_ejemplo_servicio.xlsx", "Datos de ejemplo — servicios de entrega ficticios (Cap. 6)")
    (AQUI / "cap06_ejemplo.json").write_text(json.dumps({"libro": LIBRO, "servicio": SERVICIO}, ensure_ascii=False), encoding="utf-8")
    print(f"[OK] {AQUI / 'cap06_ejemplo.json'}")


if __name__ == "__main__":
    main()
