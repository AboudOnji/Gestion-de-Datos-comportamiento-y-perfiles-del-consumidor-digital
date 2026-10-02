"""crear_excel_actividad_cap02.py — Datos de la actividad adicional del Cap. 2

Autor:       Dr. Aboud Barsekh-Onji · IPN — ESCA Unidad Santo Tomás
Descripción: Genera matlab/cap02/cap02_actividad_tlacoyo_datos.xlsx: 240 clientes
             SINTÉTICOS de «Tlacoyo Express» (app ficticia de comida a domicilio)
             para usarse con cap02_segmentacion_datos_propios.html. Misma
             dinámica que el ejemplo del capítulo: 3 variables y 3 perfiles
             ocultos con traslape deliberado; distinta historia y distintas
             variables. Los perfiles ocultos NO siguen un patrón «bajo/medio/alto»
             (quien más pide es quien menos gasta), para que la propuesta
             automática de la app no baste y los estudiantes formulen sus
             propios arquetipos a partir del contexto.
             Semilla fija (42): el archivo es reproducible.
Uso:         conda run -n research python matlab/utils/apps/datos/crear_excel_actividad_cap02.py
"""
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).resolve().parent
DESTINO = AQUI.parents[2] / "cap02" / "cap02_actividad_tlacoyo_datos.xlsx"
GUINDA = "6F1D46"

ENCABEZADOS = ["ID_cliente", "Pedidos_al_mes", "Gasto_por_pedido_MXN", "Notificaciones_abiertas_pct"]

# Perfiles OCULTOS (solo para generar los datos; no se entregan a los alumnos).
# Calibrados para reproducir el traslape del ejemplo del capítulo: silueta
# ~0.58 («razonable») y ~18 % de pertenencia ambigua.
# [pedidos al mes, gasto por pedido (MXN), % de notificaciones abiertas]
PERFILES = {
    "Oficinista entre semana": {"n": 100, "centro": [8.0, 165, 50], "disp": [2.3, 45, 14]},
    "Familia de fin de semana": {"n": 70, "centro": [3.5, 430, 32], "disp": [1.3, 120, 13]},
    "Cazador de promociones": {"n": 70, "centro": [5.5, 215, 75], "disp": [1.8, 60, 12]},
}


def generar():
    rng = np.random.default_rng(42)
    filas = []
    for p in PERFILES.values():
        z = rng.standard_normal((p["n"], 3))
        x = np.array(p["centro"]) + z * np.array(p["disp"])
        for f in x:
            filas.append([
                round(float(max(f[0], 0.5)), 1),          # pedidos al mes (promedio trimestral), > 0
                round(float(max(f[1], 60.0)), 2),         # gasto por pedido, mínimo de la app: 60 MXN
                int(round(min(max(f[2], 0), 100))),       # % de notificaciones abiertas, 0-100
            ])
    orden = rng.permutation(len(filas))
    return [[f"T{i + 1:04d}"] + filas[j] for i, j in enumerate(orden)]


def main():
    clientes = generar()
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(ENCABEZADOS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=GUINDA)
        c.alignment = Alignment(horizontal="center")
    for f in clientes:
        ws.append(f)
    for fila in ws.iter_rows(min_row=2, min_col=2, max_col=4):
        fila[0].number_format = "0.0"
        fila[1].number_format = "0.00"
        fila[2].number_format = "0"
    for col, ancho in zip("ABCD", (12, 17, 23, 29)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    wb.properties.creator = "Instituto Politécnico Nacional - Dr. Aboud Barsekh Onji"
    wb.properties.title = "Tlacoyo Express (empresa ficticia) — datos sintéticos, actividad adicional Cap. 2"
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO)
    print(f"[OK] {DESTINO} ({len(clientes)} clientes)")


if __name__ == "__main__":
    main()
