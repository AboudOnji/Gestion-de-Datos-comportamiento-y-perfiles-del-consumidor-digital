"""crear_excel_ejemplo_cap03.py — Genera matlab/cap03/cap03_datos_ejemplo.xlsx

Autor:       Dr. Aboud Barsekh-Onji · IPN — ESCA Unidad Santo Tomás
Descripción: Archivo de ejemplo (y plantilla) para la app
             cap03_rfm_datos_propios.html. Dos hojas:
             «Datos»        — 300 clientes SINTÉTICOS (los mismos que genera
                              cap03_rfm_clv_ab.m con rng(42): recencia, frecuencia
                              y monto promedio, con 5 % de montos faltantes y 2 %
                              de atípicos), con el monto redondeado a centavos
                              como lo entregaría un sistema real.
             «Experimento»  — una fila por visitante de la prueba A/B del Caso A
                              (A: 1,200 visitantes, 132 conversiones; B: 1,180,
                              159), en orden mezclado.
             También escribe datos/cap03_ejemplo.json (los mismos datos, para el
             botón «Usar datos de ejemplo» de la app).
Uso:         conda run -n research python matlab/utils/apps/datos/crear_excel_ejemplo_cap03.py
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).resolve().parent
DESTINO = AQUI.parents[2] / "cap03" / "cap03_datos_ejemplo.xlsx"
GUINDA = "6F1D46"

ENC_DATOS = ["ID_cliente", "Recencia_dias", "Frecuencia_compras", "Monto_promedio_MXN"]
ENC_EXP = ["ID_visita", "Version", "Convirtio"]


def clientes():
    base = json.loads((AQUI / "cap03_datos.json").read_text())["300"]
    filas = []
    for i, (r, f, m) in enumerate(zip(base["r"], base["f"], base["m"])):
        filas.append([f"C{i + 1:04d}", int(r), int(f), None if m is None else round(m, 2)])
    return filas


def experimento():
    nA, cA, nB, cB = 1200, 132, 1180, 159
    filas = [("A", 1 if i < cA else 0) for i in range(nA)] + [("B", 1 if i < cB else 0) for i in range(nB)]
    random.Random(42).shuffle(filas)
    return [[f"V{i + 1:05d}", v, c] for i, (v, c) in enumerate(filas)]


def encabezado(ws, nombres, anchos):
    ws.append(nombres)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=GUINDA)
        c.alignment = Alignment(horizontal="center")
    for col, ancho in zip("ABCDEF", anchos):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"


def main():
    cl, ex = clientes(), experimento()
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    encabezado(ws, ENC_DATOS, (13, 15, 19, 21))
    for f in cl:
        ws.append(f)
    for fila in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        fila[0].number_format = "0.00"
    we = wb.create_sheet("Experimento")
    encabezado(we, ENC_EXP, (12, 10, 11))
    for f in ex:
        we.append(f)
    wb.properties.creator = "Instituto Politécnico Nacional - Dr. Aboud Barsekh Onji"
    wb.properties.title = "Datos de ejemplo — RFM, CLV y prueba A/B (Cap. 3)"
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO)
    print(f"[OK] {DESTINO}")
    (AQUI / "cap03_ejemplo.json").write_text(json.dumps(
        {"datos": {"encabezados": ENC_DATOS, "filas": cl}, "experimento": {"encabezados": ENC_EXP, "filas": ex}},
        ensure_ascii=False), encoding="utf-8")
    print(f"[OK] {AQUI / 'cap03_ejemplo.json'}")


if __name__ == "__main__":
    main()
