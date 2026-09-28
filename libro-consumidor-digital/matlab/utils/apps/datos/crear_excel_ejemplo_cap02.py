"""crear_excel_ejemplo_cap02.py — Genera matlab/cap02/cap02_datos_ejemplo.xlsx

Autor:       Dr. Aboud Barsekh-Onji · IPN — ESCA Unidad Santo Tomás
Descripción: Archivo de ejemplo (y plantilla) con UNA sola hoja, «Datos»:
             solo los datos de los clientes (las instrucciones de formato y los
             arquetipos están en la app). Para la app
             cap02_segmentacion_datos_propios.html. Los datos son los 240
             clientes SINTÉTICOS del Caso A de cap02_segmentacion.m (mismo
             flujo randn de MATLAB tras rng(42)), redondeados como vendrían de
             un sistema real (frecuencia con 1 decimal, ticket con 2, engagement
             entero), acotados a rangos válidos y en orden mezclado, con un
             identificador anónimo por cliente.
             También escribe datos/cap02_ejemplo.json (los mismos datos, para
             el botón «Usar datos de ejemplo» de la app).
Uso:         conda run -n research python matlab/utils/apps/datos/crear_excel_ejemplo_cap02.py
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).resolve().parent
DESTINO = AQUI.parents[2] / "cap02" / "cap02_datos_ejemplo.xlsx"

CENTROS = [[1.5, 350, 25], [4.0, 900, 78], [2.6, 600, 52]]
DISPERSION = [[0.6, 90, 10], [0.9, 150, 12], [0.8, 120, 14]]
N_POR_SEGMENTO = 80

GUINDA = "6F1D46"


def generar_clientes():
    z = json.loads((AQUI / "randn42.json").read_text())
    filas, off, n = [], 0, N_POR_SEGMENTO
    for c, d in zip(CENTROS, DISPERSION):
        for i in range(n):
            v = [c[j] + z[off + j * n + i] * d[j] for j in range(3)]
            filas.append([
                round(max(v[0], 0.1), 1),          # frecuencia mensual > 0
                round(max(v[1], 50.0), 2),         # ticket promedio (MXN)
                int(round(min(max(v[2], 0), 100))),  # engagement 0-100
            ])
        off += 3 * n
    random.Random(42).shuffle(filas)
    return [[f"C{i + 1:04d}"] + f for i, f in enumerate(filas)]


def hoja_datos(ws, clientes):
    ws.title = "Datos"
    encabezados = ["ID_cliente", "Frecuencia_mensual", "Ticket_promedio_MXN", "Engagement_digital"]
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=GUINDA)
        c.alignment = Alignment(horizontal="center")
    for f in clientes:
        ws.append(f)
    for fila in ws.iter_rows(min_row=2, min_col=2, max_col=4):
        fila[0].number_format = "0.0"
        fila[1].number_format = "0.00"
        fila[2].number_format = "0"
    for col, ancho in zip("ABCD", (13, 21, 22, 21)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    clientes = generar_clientes()
    hoja_datos(wb.active, clientes)
    wb.properties.creator = "Instituto Politécnico Nacional - Dr. Aboud Barsekh Onji"
    wb.properties.title = "Datos de ejemplo — segmentación del consumidor digital (Cap. 2)"
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO)
    print(f"[OK] {DESTINO}")
    # Copia de los mismos datos para el botón «Usar datos de ejemplo» de la app
    encabezados = ["ID_cliente", "Frecuencia_mensual", "Ticket_promedio_MXN", "Engagement_digital"]
    (AQUI / "cap02_ejemplo.json").write_text(
        json.dumps({"encabezados": encabezados, "filas": clientes}, ensure_ascii=False), encoding="utf-8")
    print(f"[OK] {AQUI / 'cap02_ejemplo.json'}")


if __name__ == "__main__":
    main()
