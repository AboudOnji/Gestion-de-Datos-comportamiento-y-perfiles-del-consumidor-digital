"""crear_excel_ejemplo_cap07.py — Genera los Excel de ejemplo de la app «datos propios» del Cap. 7

Autor:       Dr. Aboud Barsekh-Onji · IPN — ESCA Unidad Santo Tomás
Descripción: Dos libros de ejemplo (y plantilla) para cap07_journey_datos_propios.html:
             «cap07_datos_ejemplo.xlsx»        — 3 000 recorridos SINTÉTICOS simulados con la cadena de
                                                  canales de cap07_journey_markov.m (valores por defecto),
                                                  uno por fila: ID, Canales (Redes>Búsqueda>…), Convirtio.
             «cap07_datos_ejemplo_tienda.xlsx» — recorridos agrupados con la columna Cantidad, de una
                                                  tienda ficticia de ropa con cinco canales.
             También escribe datos/cap07_ejemplo.json para los botones de ejemplo de la app.
Uso:         conda run -n research python matlab/utils/apps/datos/crear_excel_ejemplo_cap07.py
"""
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).resolve().parent
CARPETA = AQUI.parents[2] / "cap07"
GUINDA = "6F1D46"
CANALES = ["Búsqueda", "Redes sociales", "Correo", "Directo"]
QC = np.array([[0.10, 0.05, 0.10, 0.15], [0.30, 0.15, 0.05, 0.05], [0.15, 0.05, 0.05, 0.25], [0.05, 0.00, 0.05, 0.10]])
RC = np.array([0.30, 0.02, 0.10, 0.35])
PI0 = np.array([0.35, 0.40, 0.10, 0.15])
PI0 = np.array([PI0[0], 0.40, PI0[2], PI0[3]])
resto = np.array([0.35, 0.10, 0.15]) / 0.60 * 0.60
PI0 = np.array([resto[0], 0.40, resto[1], resto[2]])


def simular(n=3000, semilla=42):
    rng = np.random.RandomState(semilla)
    filas = []
    for i in range(n):
        c = rng.choice(4, p=PI0 / PI0.sum()); ruta = [c]
        while True:
            u = rng.rand()
            acum = np.cumsum(np.concatenate([QC[c], [RC[c]]]))
            j = int(np.searchsorted(acum, u, side="right"))
            if j < 4:
                c = j; ruta.append(c)
                if len(ruta) >= 25:
                    ok = 0; break
            elif j == 4:
                ok = 1; break
            else:
                ok = 0; break
        filas.append((i + 1, ">".join(CANALES[x] for x in ruta), ok))
    return filas


TIENDA = [  # (ruta, convirtio, cantidad) — tienda ficticia de ropa
    ("Redes>Búsqueda>Directo", "Sí", 160), ("Redes>Búsqueda", "No", 240), ("Redes", "No", 520), ("Búsqueda>Directo", "Sí", 210),
    ("Búsqueda", "No", 380), ("Búsqueda", "Sí", 90), ("Correo>Directo", "Sí", 130), ("Correo", "No", 260), ("Directo", "Sí", 175),
    ("Directo", "No", 150), ("Marketplace>Directo", "Sí", 60), ("Marketplace", "No", 310), ("Marketplace", "Sí", 45),
    ("Redes>Correo>Directo", "Sí", 55), ("Redes>Correo", "No", 120), ("Redes>Redes>Búsqueda>Directo", "Sí", 70), ("Redes>Redes>Búsqueda", "No", 95),
    ("Búsqueda>Marketplace", "No", 85), ("Búsqueda>Marketplace>Directo", "Sí", 40), ("Redes>Marketplace", "No", 110),
    ("Búsqueda>Correo>Directo", "Sí", 50), ("Búsqueda>Búsqueda>Directo", "Sí", 65), ("Correo>Búsqueda", "No", 70), ("Redes>Directo", "Sí", 35),
]


def estilo(ws, anchos):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=GUINDA); c.alignment = Alignment(horizontal="center")
    for col, ancho in zip("ABCD", anchos):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"


def libro(encabezados, filas, anchos, ruta, titulo):
    wb = Workbook(); ws = wb.active; ws.title = "Recorridos"
    ws.append(encabezados)
    for f in filas:
        ws.append(list(f))
    estilo(ws, anchos)
    wb.properties.creator = "Instituto Politécnico Nacional - Dr. Aboud Barsekh Onji"; wb.properties.title = titulo
    CARPETA.mkdir(parents=True, exist_ok=True)
    wb.save(ruta); print(f"[OK] {ruta}")


def main():
    s1 = simular()
    libro(["ID", "Canales", "Convirtio"], s1, (8, 60, 12), CARPETA / "cap07_datos_ejemplo.xlsx", "Datos de ejemplo — 3 000 recorridos sintéticos (Cap. 7)")
    libro(["Ruta", "Convirtio", "Cantidad"], TIENDA, (40, 12, 10), CARPETA / "cap07_datos_ejemplo_tienda.xlsx", "Datos de ejemplo — tienda ficticia de ropa (Cap. 7)")
    (AQUI / "cap07_ejemplo.json").write_text(json.dumps({
        "libro": {"encabezados": ["ID", "Canales", "Convirtio"], "filas": [list(f) for f in s1]},
        "tienda": {"encabezados": ["Ruta", "Convirtio", "Cantidad"], "filas": [list(f) for f in TIENDA]}}, ensure_ascii=False), encoding="utf-8")
    print(f"[OK] {AQUI / 'cap07_ejemplo.json'}  (conversión del ejemplo 1: {np.mean([f[2] for f in s1]):.3f})")


if __name__ == "__main__":
    main()
